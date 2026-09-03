# -*- coding: utf-8 -*-
"""
阶段 6：终稿 Markdown → 规范 Word 文档（替代旧 md_to_word.py）
修复/增强：
- 自动跳过 YAML front matter 与 <sup>/</sup> HTML 标签（旧版会原样显示在 Word 中）
- 正文 [SE01] 按参考文献出现顺序映射为 [1] 并设为上标；缺失映射时告警
- 支持引用行(>)、无序列表(-/*)、分割线；中西文字体/字号由 project.yaml 驱动
用法：
  python -m workflow.s6_to_word --config config/project.yaml
  python -m workflow.s6_to_word --config config/project.yaml -i in.md -o out.docx
"""
from __future__ import annotations

import argparse
import re
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from .common import WorkflowConfig, split_frontmatter, strip_sup_tags


class WordStyle:
    def __init__(self, cfg: WorkflowConfig):
        w = cfg.word_cfg
        self.font_cn = w.get("font_cn", "仿宋")
        self.font_en = w.get("font_en", "Times New Roman")
        self.size_title = Pt(w.get("size_title_pt", 22))
        self.size_heading = Pt(w.get("size_heading_pt", 16))
        self.size_body = Pt(w.get("size_body_pt", 16))
        self.size_cite = Pt(w.get("size_citation_pt", 12))
        self.line_spacing = w.get("line_spacing", 1.5)
        self.prefix = cfg.id_prefix


def set_run_font(run, style: WordStyle, size=None, bold=False):
    run.font.name = style.font_en
    run.font.size = size or style.size_body
    run.font.bold = bold
    rpr = run._r.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), style.font_en)
    rfonts.set(qn("w:hAnsi"), style.font_en)
    rfonts.set(qn("w:eastAsia"), style.font_cn)
    rfonts.set(qn("w:cs"), style.font_cn)


def para_format(p, *, indent_cm=None, before=0, after=6, spacing=1.5, align=None):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = spacing
    if indent_cm is not None:
        pf.first_line_indent = Cm(indent_cm)
    if align is not None:
        p.alignment = align


def configure_document(doc: Document, style: WordStyle):
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Cm(2.54)
    sec.left_margin = sec.right_margin = Cm(3.0)
    normal = doc.styles["Normal"]
    normal.font.name = style.font_en
    normal.font.size = style.size_body
    rpr = normal._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for k, v in (("w:ascii", style.font_en), ("w:hAnsi", style.font_en),
                 ("w:eastAsia", style.font_cn), ("w:cs", style.font_cn)):
        rfonts.set(qn(k), v)
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        st = doc.styles[name]
        st.font.name = style.font_en
        st.font.size = style.size_heading
        st.font.bold = True


def add_hyperlink(paragraph, text, url, style: WordStyle):
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rf = OxmlElement("w:rFonts")
    for k, v in (("w:ascii", style.font_en), ("w:hAnsi", style.font_en),
                 ("w:eastAsia", style.font_cn), ("w:cs", style.font_cn)):
        rf.set(qn(k), v)
    rpr.append(rf)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(style.size_body.pt * 2)))
    rpr.append(sz)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(u)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


TOKEN_RE = re.compile(r"\*\*([^*]+?)\*\*|`([^`]+)`|\[([^\]]+)\]\(([^)]+)\)")


def add_inline(paragraph, text: str, style: WordStyle, size=None):
    pos = 0
    for m in TOKEN_RE.finditer(text):
        if m.start() > pos:
            r = paragraph.add_run(text[pos:m.start()])
            set_run_font(r, style, size=size)
        if m.group(1) is not None:
            r = paragraph.add_run(m.group(1))
            set_run_font(r, style, size=size, bold=True)
        elif m.group(2) is not None:
            r = paragraph.add_run(m.group(2))
            set_run_font(r, style, size=size)
        else:
            add_hyperlink(paragraph, m.group(3), m.group(4), style)
        pos = m.end()
    if pos < len(text):
        r = paragraph.add_run(text[pos:])
        set_run_font(r, style, size=size)


def add_cited_text(paragraph, text: str, ref_map: dict, style: WordStyle, size=None):
    """文本中 [SE01] → [n] 上标；其余走行内格式。"""
    cite_re = re.compile(rf"\[{style.prefix}\d+\]")
    pos = 0
    for m in cite_re.finditer(text):
        if m.start() > pos:
            add_inline(paragraph, text[pos:m.start()], style, size)
        sid = m.group(0).strip("[]")
        if sid in ref_map:
            visible = f"[{ref_map[sid]}]"
        else:
            visible = f"[{sid}]"
            print(f"  ⚠️ 正文引用 {sid} 未在参考文献中找到编号映射，保留原编号")
        r = paragraph.add_run(visible)
        set_run_font(r, style, size=style.size_cite)
        r.font.superscript = True
        pos = m.end()
    if pos < len(text):
        add_inline(paragraph, text[pos:], style, size)


def build_ref_map(lines: Iterable[str], prefix: str) -> dict:
    in_refs, mapping, n = False, {}, 1
    ref_re = re.compile(rf"^\[{prefix}(\d+)\]\s*(.*)$")
    for raw in lines:
        s = raw.strip()
        if s.startswith("#") and "参考文献" in s:
            in_refs = True
            continue
        if not in_refs:
            continue
        m = ref_re.match(s)
        if m:
            sid = f"{prefix}{int(m.group(1)):02d}"
            if sid not in mapping:
                mapping[sid] = n
                n += 1
    return mapping


def convert(md_path: Path, docx_path: Path, cfg: WorkflowConfig):
    style = WordStyle(cfg)
    raw = md_path.read_text(encoding="utf-8")
    raw = split_frontmatter(raw)[1] if raw.lstrip().startswith("---") else raw
    raw = strip_sup_tags(raw)  # HTML 上角标标签由 docx 重新生成
    lines = raw.splitlines()
    ref_map = build_ref_map(lines, cfg.id_prefix)
    if not ref_map:
        raise ValueError("未找到“## 参考文献”及 [SE01] 形态的条目，无法建立引用编号映射")
    print(f"📚 参考文献编号映射 {len(ref_map)} 条")

    doc = Document()
    configure_document(doc, style)
    in_refs = False
    ref_line_re = re.compile(rf"^\[{cfg.id_prefix}\d+\]\s*(.*)$")

    for raw_line in lines:
        s = raw_line.rstrip()
        stripped = s.strip()
        if not stripped or stripped == "---":
            continue
        # 主标题
        if stripped.startswith("# "):
            p = doc.add_paragraph()
            para_format(p, before=0, after=18, spacing=1.2,
                        align=WD_ALIGN_PARAGRAPH.CENTER)
            r = p.add_run(stripped[2:].strip())
            set_run_font(r, style, size=style.size_title, bold=True)
            continue
        if stripped.startswith("## "):
            h = stripped[3:].strip()
            in_refs = h == "参考文献"
            p = doc.add_paragraph()
            para_format(p, before=10, after=6, spacing=1.2, align=WD_ALIGN_PARAGRAPH.LEFT)
            r = p.add_run(h)
            set_run_font(r, style, size=style.size_heading, bold=True)
            continue
        if stripped.startswith("### "):
            p = doc.add_paragraph()
            para_format(p, before=8, after=6, spacing=1.2, align=WD_ALIGN_PARAGRAPH.LEFT)
            r = p.add_run(stripped[4:].strip())
            set_run_font(r, style, size=style.size_heading, bold=True)
            continue
        # 参考文献条目：[SE01] → [n]（非上标），其余正常排版
        if in_refs:
            m = ref_line_re.match(stripped)
            if m:
                sid = re.match(rf"\[{cfg.id_prefix}\d+\]", stripped).group(0).strip("[]")
                num = ref_map.get(sid)
                p = doc.add_paragraph()
                para_format(p, after=6, spacing=style.line_spacing,
                            align=WD_ALIGN_PARAGRAPH.LEFT)
                r = p.add_run(f"[{num}]" if num else f"[{sid}]")
                set_run_font(r, style, size=style.size_body)
                add_inline(p, " " + m.group(1), style, size=style.size_body)
                continue
        # 引用行
        if stripped.startswith(">"):
            stripped = stripped.lstrip("> ").strip()
        # 无序列表
        bullet = bool(re.match(r"^[-*]\s+", stripped))
        if bullet:
            stripped = "• " + re.sub(r"^[-*]\s+", "", stripped)
        # 关键词加粗引导
        if stripped.startswith("**关键词**"):
            p = doc.add_paragraph()
            para_format(p, after=10, spacing=style.line_spacing,
                        align=WD_ALIGN_PARAGRAPH.JUSTIFY)
            r = p.add_run("关键词")
            set_run_font(r, style, bold=True)
            rest = re.sub(r"^\*\*关键词\*\*", "", stripped).lstrip("：:")
            add_cited_text(p, rest, ref_map, style)
            continue

        p = doc.add_paragraph()
        para_format(p, indent_cm=None if bullet else 0.74,
                    after=6, spacing=style.line_spacing,
                    align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        add_cited_text(p, stripped, ref_map, style)

    props = doc.core_properties
    props.title = cfg.report_title
    props.comments = "由 research_report_workflow 阶段6自动生成"
    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(docx_path)
    print(f"✅ Word 已生成：{docx_path}")


def run(cfg: WorkflowConfig, input_path: Path | None = None, output_path: Path | None = None):
    md = input_path or cfg.final_md_path
    out = output_path or cfg.final_docx_path
    if not md.exists():
        raise FileNotFoundError(f"终稿 Markdown 不存在：{md}（请先运行阶段5）")
    print("=" * 72)
    print("🚀 阶段6：Markdown → Word")
    print("=" * 72)
    convert(md, out, cfg)
    return {"path": str(out)}


def main():
    ap = argparse.ArgumentParser(description="Markdown 调研报告转 Word")
    ap.add_argument("--config", required=True)
    ap.add_argument("-i", "--input")
    ap.add_argument("-o", "--output")
    ap.add_argument("--root")
    args = ap.parse_args()
    from .common import load_config
    cfg = load_config(args.config, args.root)
    run(cfg,
        Path(args.input).resolve() if args.input else None,
        Path(args.output).resolve() if args.output else None)


if __name__ == "__main__":
    main()

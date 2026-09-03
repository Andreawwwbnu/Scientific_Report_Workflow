# -*- coding: utf-8 -*-
"""
离线自测：不依赖网络与 LLM。
构造一份贴近阶段4输出的合成草稿（故意混入：
裸引用、<sup>引用、镜像编号 SE06、YAML 头缺失场景），
跑阶段5、阶段6，并回读 docx 验证：
  1. Word 中不出现字面量 <sup>/</sup>/---
  2. 正文 [SE01] 被映射为数字上标 [1]
  3. 参考文献 13 条（15 条清单去掉 2 个镜像）
  4. 镜像 SE06 归并为 SE01
"""
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from workflow.common import load_config
from workflow import s4_draft, s5_finalize, s6_to_word

cfg = load_config(ROOT / "config" / "project.yaml")
refs = s4_draft.build_references(cfg)
assert len(refs.splitlines()) == 13, f"主文献应为13条，实际 {len(refs.splitlines())}"

para = ("自进化智能体区别于静态指令遵循系统的核心，在于其经验依赖、持久记忆与自主探索三项特性，"
        "使其能够在部署后持续改写自身组件[SE01]。在工程诱因层面，文本梯度与多提示优化把提示、"
        "拓扑与记忆纳入统一的动态优化空间[SE03]。")
para2 = ("系统通过输入、智能体、环境与优化器构成的四部式闭环持续更新<sup>[SE05]</sup>，"
         "而演化对象、时机与手段构成正交三维度[SE06]。安全适应与性能保持两条定律构成其合规边界[SE04]。")
para3 = ("源码级自改写标志着从文本工件调优向图灵完备自覆写的跨越[SE08]；多智能体协同演化与"
         "可执行子智能体积累代表两条互补路线[SE09][SE10]。开源生态降低了工程化门槛[SE11]，"
         "但商业化仍处早期阶段[SE12]。在运筹优化等深水区任务中，图介导的路径重组提供了可解释"
         "落地范式[SE13]，上下文密度压缩显著改善长程任务的 Token 经济性[SE14]，企业治理框架下的"
         "人类介入阻断机制则界定了可控自进化的部署红线[SE15]。演进价值与投入产出比在产业评论中"
         "亦得到论证[SE02]。")

draft = f"""# {cfg.report_title}

## 摘要
本报告围绕智能体自进化框架，系统梳理其遴选价值、技术内涵、最新进展与应用前景，基于13篇唯一文献形成证据链[SE01]。

## 关键词
智能体；自进化；多智能体；记忆演化；源码级改写；治理

## 一、遴选理由
### 1.1 核心技术特征与代际优势
{para}
### 1.2 产业价值与突破诱因
{para3}

## 二、技术内涵
### 2.1 科学定义与闭环运行机理
{para2}
### 2.2 演化维度与技术边界
{para2}

## 三、最新进展
### 3.1 技术演进与里程碑成果
{para3}
### 3.2 开源生态与产业动态
{para3}

## 四、应用前景与重要价值
### 4.1 典型落地应用场景
{para3}
### 4.2 多维度战略价值
{para3}

## 结论与展望
自进化框架正沿提示层、代码级与多智能体协同三条路线收敛，工程化与治理将共同决定其落地节奏。

## 参考文献

{refs}
"""
cfg.draft_dir.mkdir(parents=True, exist_ok=True)
cfg.draft_path.write_text(draft, encoding="utf-8")
print("合成草稿已写入：", cfg.draft_path)

r5 = s5_finalize.run(cfg)
r6 = s6_to_word.run(cfg)

# ---------- 回读验证 ----------
from docx import Document
doc = Document(str(cfg.final_docx_path))
all_text = "\n".join(p.text for p in doc.paragraphs)

problems = []
for leak in ("<sup>", "</sup>", "---", "SE06", "SE07"):
    if leak in all_text:
        problems.append(f"Word 中出现残留：{leak}")
if "[SE" in all_text:
    # 参考文献条目允许 SE？不——参考文献也已转数字编号，因此不应残留 SE
    problems.append("Word 中仍存在 SE 形态编号")
if "[1]" not in all_text:
    problems.append("未见数字角标 [1]")
# 上标 run 检查
sup_runs = [r.text for p in doc.paragraphs for r in p.runs if r.font.superscript]
if not sup_runs:
    problems.append("未检测到任何上标 run")

print("\n================ 自测结果 ================")
print("终稿质检 errors:", r5["errors"])
print("正文字数:", r5["lengths"])
print("上标角标示例:", sorted(set(sup_runs))[:8], "共", len(sup_runs), "处")
print("Word 段落数:", len(doc.paragraphs))
if problems:
    print("❌ 发现问题：")
    for p in problems:
        print("  -", p)
    sys.exit(1)
print("✅ 全部断言通过：无标签/YAML残留，编号映射与上标正常，镜像已归并")

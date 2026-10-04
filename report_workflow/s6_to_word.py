# -*- coding: utf-8 -*-
"""阶段 6：终稿 Markdown → 规范 Word 文档。

v3 相对旧版：
- 标题使用 Word 内置 Title / Heading 1 / Heading 2 样式（有导航窗格、可生成目录），
  字体字号集中在样式里设置（旧版全部是普通段落）；
- 可选：目录域（word.toc: true，打开 Word 时提示更新域）、页码页脚、
  自定义模板（word.template: 指向含你单位样式的 .docx/.dotx）；
- 摘要/关键词既可用标题（默认），也可用行内标签（word.abstract_style: inline）；
- 正文 [LBQ01] 按参考文献顺序映射为上标 [n]；缺失映射时告警；参考文献悬挂缩进、URL 可点击。
"""
from __future__ import annotations

import re
from pathlib import Path
from typing import Dict, List, Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .config import WorkflowConfig
from .textutil import split_frontmatter, strip_sup_tags

_THEME_ATTRS = ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme")
URL_RE = re.compile(r"(https?://[^\s\]）)，。；]+)")


class WordStyle:
    def __init__(self, cfg: WorkflowConfig):
        w = cfg.word_cfg
        self.font_cn = w.get("font_cn", "仿宋")
        self.font_en = w.get("font_en", "Times New Roman")
        self.size_title = Pt(w.get("size_title_pt", 22))
        self.size_heading = Pt(w.get("size_heading_pt", 16))
        self.size_body = Pt(w.get("size_body_pt", 16))
        self.size_cite = Pt(w.get("size_citation_pt", 12))
        self.line_spacing = w.get("line_spacing", 1.5)
        self.toc = bool(w.get("toc", False))
        self.page_numbers = bool(w.get("page_numbers", True))
        self.abstract_style = w.get("abstract_style", "heading")
        self.template = w.get("template")
        self.prefix = cfg.id_prefix
        self.indent = Pt(self.size_body.pt * 2)    # 首行缩进 2 字符


# ---------------- 字体 ----------------
def _set_rfonts(rpr, style: WordStyle):
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for a in _THEME_ATTRS:
        rfonts.attrib.pop(qn(f"w:{a}"), None)
    for k, v in (("w:ascii", style.font_en), ("w:hAnsi", style.font_en),
                 ("w:eastAsia", style.font_cn), ("w:cs", style.font_cn)):
        rfonts.set(qn(k), v)


def set_run_font(run, style: WordStyle, size=None, bold=False):
    run.font.name = style.font_en
    run.font.size = size or style.size_body
    run.font.bold = bold
    _set_rfonts(run._r.get_or_add_rPr(), style)


def configure_style(st, style: WordStyle, size, bold=None, align=None, before=0, after=6, black=True):
    st.font.name = style.font_en
    st.font.size = size
    if bold is not None:
        st.font.bold = bold
    if black:
        st.font.color.rgb = RGBColor(0, 0, 0)
    st.font.italic = False
    _set_rfonts(st.element.get_or_add_rPr(), style)
    pf = st.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.line_spacing = style.line_spacing
    if align is not None:
        pf.alignment = align


def new_document(style: WordStyle, project_dir: Path) -> Document:
    if style.template:
        tp = Path(style.template).expanduser()
        tp = tp if tp.is_absolute() else project_dir / tp
        doc = Document(str(tp))
        body = doc.element.body
        for child in list(body):
            if child.tag != qn("w:sectPr"):
                body.remove(child)
    else:
        doc = Document()
        sec = doc.sections[0]
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.top_margin = sec.bottom_margin = Cm(2.54)
        sec.left_margin = sec.right_margin = Cm(3.0)
    st = doc.styles
    configure_style(st["Normal"], style, style.size_body, black=False)
    configure_style(st["Title"], style, style.size_title, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=12)
    configure_style(st["Heading 1"], style, style.size_heading, bold=True, before=12, after=6)
    configure_style(st["Heading 2"], style, style.size_heading, bold=True, before=6, after=4)
    try:
        configure_style(st["List Bullet"], style, style.size_body, black=False)
    except KeyError:
        pass
    # 去掉默认 Title 样式的下边框
    ppr = st["Title"].element.get_or_add_pPr()
    for b in ppr.findall(qn("w:pBdr")):
        ppr.remove(b)
    return doc


# ---------------- 域：页码 / 目录 ----------------
def _field(paragraph, instr: str, placeholder: str = ""):
    def fld(t):
        r = OxmlElement("w:r")
        c = OxmlElement("w:fldChar")
        c.set(qn("w:fldCharType"), t)
        r.append(c)
        return r

    paragraph._p.append(fld("begin"))
    r = OxmlElement("w:r")
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = f" {instr} "
    r.append(it)
    paragraph._p.append(r)
    paragraph._p.append(fld("separate"))
    if placeholder:
        rt = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.text = placeholder
        rt.append(t)
        paragraph._p.append(rt)
    paragraph._p.append(fld("end"))


def add_page_numbers(doc: Document, style: WordStyle):
    for sec in doc.sections:
        p = sec.footer.paragraphs[0] if sec.footer.paragraphs else sec.footer.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _field(p, "PAGE", "1")
        for r in p.runs:
            set_run_font(r, style, size=Pt(10.5))


def add_toc(doc: Document, style: WordStyle):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("目 录")
    set_run_font(r, style, size=style.size_heading, bold=True)
    p2 = doc.add_paragraph()
    _field(p2, 'TOC \\o "1-2" \\h \\z \\u', "（打开文档后右键此处 → 更新域，生成目录）")
    doc.add_page_break()
    settings = doc.settings.element
    if settings.find(qn("w:updateFields")) is None:
        uf = OxmlElement("w:updateFields")
        uf.set(qn("w:val"), "true")
        settings.append(uf)


# ---------------- 行内内容 ----------------
def add_hyperlink(paragraph, text: str, url: str, style: WordStyle):
    r_id = paragraph.part.relate_to(
        url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    _set_rfonts(rpr, style)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(style.size_body.pt * 2)))
    rpr.append(sz)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(u)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


TOKEN_RE = re.compile(r"\*\*([^*]+?)\*\*|`([^`]+)`|\[([^\]]+)\]\((https?://[^)]+)\)")


def add_inline(paragraph, text: str, style: WordStyle, size=None):
    pos = 0
    for m in TOKEN_RE.finditer(text):
        if m.start() > pos:
            set_run_font(paragraph.add_run(text[pos:m.start()]), style, size=size)
        if m.group(1) is not None:
            set_run_font(paragraph.add_run(m.group(1)), style, size=size, bold=True)
        elif m.group(2) is not None:
            set_run_font(paragraph.add_run(m.group(2)), style, size=size)
        else:
            add_hyperlink(paragraph, m.group(3), m.group(4), style)
        pos = m.end()
    if pos < len(text):
        set_run_font(paragraph.add_run(text[pos:]), style, size=size)


def add_cited_text(paragraph, text: str, ref_map: Dict[str, int], style: WordStyle, size=None) -> List[str]:
    """[LBQ01] → 上标 [n]。返回未映射的编号列表。"""
    cre = re.compile(rf"\[{re.escape(style.prefix)}\d+\]")
    missing, pos = [], 0
    for m in cre.finditer(text):
        if m.start() > pos:
            add_inline(paragraph, text[pos:m.start()], style, size)
        sid = m.group(0).strip("[]")
        if sid in ref_map:
            visible = f"[{ref_map[sid]}]"
        else:
            visible = f"[{sid}]"
            missing.append(sid)
        r = paragraph.add_run(visible)
        set_run_font(r, style, size=style.size_cite)
        r.font.superscript = True
        pos = m.end()
    if pos < len(text):
        add_inline(paragraph, text[pos:], style, size)
    return missing


def build_ref_map(lines: List[str], prefix: str) -> Dict[str, int]:
    in_refs, mapping, n = False, {}, 1
    ref_re = re.compile(rf"^\[{re.escape(prefix)}(\d+)\]\s*(.*)$")
    for raw in lines:
        s = raw.strip()
        if s.startswith("#") and "参考文献" in s:
            in_refs = True
            continue
        if not in_refs:
            continue
        m = ref_re.match(s)
        if m:
            sid = f"{prefix}{m.group(1)}"
            if sid not in mapping:
                mapping[sid] = n
                n += 1
    return mapping


def _para(doc, style_name: Optional[str] = None):
    return doc.add_paragraph(style=style_name) if style_name else doc.add_paragraph()


def _fmt(p, style: WordStyle, indent=True, after=6, align=None, left_cm=None, hanging_pt=None):
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(0), Pt(after)
    pf.line_spacing = style.line_spacing
    if indent:
        pf.first_line_indent = style.indent
    if left_cm is not None:
        pf.left_indent = Cm(left_cm)
    if hanging_pt is not None:
        pf.left_indent = Pt(hanging_pt)
        pf.first_line_indent = Pt(-hanging_pt)
    if align is not None:
        p.alignment = align


def convert(md_path: Path, out_path: Path, cfg: WorkflowConfig) -> dict:
    style = WordStyle(cfg)
    doc = new_document(style, cfg.project_dir)
    _meta, body = split_frontmatter(md_path.read_text(encoding="utf-8"))
    body = strip_sup_tags(body)
    lines = body.splitlines()
    ref_map = build_ref_map(lines, style.prefix)
    ref_re = re.compile(rf"^\[{re.escape(style.prefix)}(\d+)\]\s*(.*)$")

    missing_all: List[str] = []
    in_refs = False
    toc_done = not style.toc
    n_headings = n_paras = 0
    pending_label: Optional[str] = None

    for raw in lines:
        s = raw.rstrip()
        if not s.strip() or re.fullmatch(r"-{3,}|\*{3,}", s.strip()):
            continue
        if s.startswith("# "):
            p = _para(doc, "Title")
            set_run_font(p.add_run(s[2:].strip()), style, size=style.size_title, bold=True)
            if not toc_done:
                add_toc(doc, style)
                toc_done = True
            continue
        if s.startswith("## "):
            title = s[3:].strip()
            in_refs = "参考文献" in title
            if style.abstract_style == "inline" and title in ("摘要", "关键词"):
                pending_label = title
                continue
            p = _para(doc, "Heading 1")
            set_run_font(p.add_run(title), style, size=style.size_heading, bold=True)
            n_headings += 1
            continue
        if s.startswith("### "):
            p = _para(doc, "Heading 2")
            set_run_font(p.add_run(s[4:].strip()), style, size=style.size_heading, bold=True)
            n_headings += 1
            continue
        if s.startswith("#### "):
            p = _para(doc)
            set_run_font(p.add_run(s[5:].strip()), style, bold=True)
            _fmt(p, style, indent=False)
            continue
        if in_refs:
            m = ref_re.match(s.strip())
            p = _para(doc)
            if m:
                sid = f"{style.prefix}{m.group(1)}"
                n = ref_map.get(sid, "?")
                set_run_font(p.add_run(f"[{n}] "), style)
                rest = m.group(2)
                pos = 0
                for um in URL_RE.finditer(rest):
                    if um.start() > pos:
                        set_run_font(p.add_run(rest[pos:um.start()]), style)
                    add_hyperlink(p, um.group(1), um.group(1), style)
                    pos = um.end()
                if pos < len(rest):
                    set_run_font(p.add_run(rest[pos:]), style)
            else:
                add_inline(p, s.strip(), style)
            _fmt(p, style, indent=False, after=3, hanging_pt=style.size_body.pt * 2.2)
            continue
        if s.lstrip().startswith("> "):
            p = _para(doc)
            add_cited_text(p, s.lstrip()[2:], ref_map, style)
            _fmt(p, style, indent=False, left_cm=1.0)
            continue
        if re.match(r"^\s*[-*]\s+", s):
            try:
                p = _para(doc, "List Bullet")
            except KeyError:
                p = _para(doc)
            missing_all += add_cited_text(p, re.sub(r"^\s*[-*]\s+", "", s), ref_map, style)
            _fmt(p, style, indent=False, after=3)
            continue
        p = _para(doc)
        if pending_label:
            set_run_font(p.add_run(f"{pending_label}："), style, bold=True)
            pending_label = None
        missing_all += add_cited_text(p, s.strip(), ref_map, style)
        _fmt(p, style, indent=True)
        n_paras += 1

    if style.page_numbers:
        add_page_numbers(doc, style)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    missing_all = sorted(set(missing_all))
    for sid in missing_all:
        print(f"  ⚠️ 正文引用 {sid} 未在参考文献中找到编号映射，保留原编号")
    print(f"📄 Word：{out_path}（{n_headings} 个标题、{n_paras} 个正文段、{len(ref_map)} 条参考文献）")
    return {"path": str(out_path), "headings": n_headings, "paragraphs": n_paras,
            "references": len(ref_map), "unmapped": missing_all}


def run(cfg: WorkflowConfig, input_path: Optional[Path] = None, output_path: Optional[Path] = None, **_) -> dict:
    md = input_path or cfg.final_md_path
    out = output_path or cfg.final_docx_path
    print("=" * 72)
    print("🚀 阶段6：Markdown → Word")
    print("=" * 72)
    if not md.exists():
        print(f"❌ 终稿 Markdown 不存在：{md}（请先运行阶段5）")
        return {"ok": False, "reason": "missing_final"}
    res = convert(md, out, cfg)
    res["ok"] = not res["unmapped"]
    return res
